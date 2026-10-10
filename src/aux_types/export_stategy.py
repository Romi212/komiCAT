from abc import ABC, abstractmethod
from dataclasses import dataclass, field

from aux_types.segment_types import SegmentType
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment

@dataclass
class Row:
    kind: str                      # "band" (black separator row) or "segment"
    cells: list = field(default_factory=list)
    style: str = ""    

class Exporter(ABC):
    name = ""            # shown in the UI
    extension = ""       # "txt", "xlsx"

    @abstractmethod
    def build(self, chapter): ...

    @abstractmethod
    def export(self, data, path): ...

EXPORTERS = {}

def register(cls):
    EXPORTERS[cls.name] = cls()
    return cls


@register
class PlainTextExporter(Exporter):
    name = "Plain text"
    extension = "txt"
    combined_separator = " // "
    page_separator = "--------------------<page_name>--------------------"
    segment_prefix = ""
    segment_separator = "\n"
    bubble_text = "<text>"
    free_text = "[<text>]"
    sfx_text = "[SFX: <text>]"

    @property
    def templates(self):
        return {
            SegmentType.BUBBLE: self.bubble_text,
            SegmentType.FREE: self.free_text,      # use whatever your enum member is called
            SegmentType.SFX: self.sfx_text,
        }

    def format_segment(self, seg):
        template = self.templates.get(seg.segment_type, self.bubble_text)
        return template.replace("<text>", seg.translation)

    def format_page_header(self, page):
        return self.page_separator.replace("<page_name>", page.page_name)

    def build(self, chapter):
        lines = []
        for page in chapter.pages:
            lines.append(self.format_page_header(page))
            for head in page.segments.heads:
                lines.append(self.segment_prefix + self.combined_separator.join(self.format_segment(s) for s in head.chain()))
                lines.append("")
        
        return self.segment_separator.join(lines) if self.segment_separator else "\n".join(lines)

    def export(self, text, path):
        with open(path, "w", encoding="utf-8") as f:
            f.write(text)

@register
class ExcelExporter(Exporter):
    name = "Excel (structured)"
    extension = "xlsx"
    COLUMNS = ["Japanese", "English", "SFX/Dialogue/Box/etc", "note"]
    STYLE_FILLS = {"sfx": "FFF2CC"}                  # row style -> hex fill

    def build(self, chapter):
        rows = [Row("band", self.COLUMNS),
                Row("band", ["Title"]),
                Row("segment", [chapter.title_jp, chapter.title_tl, "Title"]),
                Row("band", ["Chapter name"]),
                Row("segment", [chapter.name_jp, chapter.name_tl, "Chapter name"])]
        for page in chapter.pages:
            rows.append(Row("band", [f"Page {page.number + 1}"]))
            for head in page.segments.heads:
                segs = list(head.chain())
                rows.append(Row("segment", [
                    "\n".join(s.source_text for s in segs),
                    "\n".join(s.translation for s in segs),
                    head.segment_type.label,
                    getattr(head, "note", ""),
                ], style=head.segment_type.key))
        return rows

    def export(self, rows, path):
        wb = Workbook()
        ws = wb.active
        ncols = len(self.COLUMNS)
        for r in rows:
            ws.append(r.cells)
            for col in range(1, ncols + 1):
                cell = ws.cell(row=ws.max_row, column=col)
                if r.kind == "band":
                    cell.fill = PatternFill("solid", start_color="000000")
                    cell.font = Font(color="FFFFFF", bold=True)
                else:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                    if r.style in self.STYLE_FILLS:
                        cell.fill = PatternFill("solid", start_color=self.STYLE_FILLS[r.style])
        for col, w in zip("ABCD", (45, 55, 20, 40)):
            ws.column_dimensions[col].width = w
        wb.save(path)